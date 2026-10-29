from pathlib import Path
from paths import BASE_DIR, resolve_path
import pandas as pd
import json
import hashlib
import os
import shutil
try:
    import pytesseract
except ImportError:
    pytesseract = None

import pymupdf
from openpyxl import load_workbook
from docx import Document
from pptx import Presentation
from PIL import Image

# 1. 파일 경로 및 기본 설정
METADATA_CSV=BASE_DIR / "metadata.csv"
OUTPUT_CSV=BASE_DIR / "extracted_text.csv"

# 한 파일에서 추출할 최대 글자 수
MAX_TEXT_CHARS = 100_000

# CSV와 엑셀은 우선 최대 100행까지만 추출
MAX_TABLE_ROWS = 100

# PDF는 우선 최대 100페이지까지만 추출
MAX_PDF_PAGES = 100

# Use PATH or an explicit override; retain the standard Windows installation.
OCR_COMMAND = os.environ.get("TESSERACT_CMD") or shutil.which("tesseract")
if not OCR_COMMAND and Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe").is_file():
    OCR_COMMAND = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
if pytesseract is not None and OCR_COMMAND:
    pytesseract.pytesseract.tesseract_cmd = OCR_COMMAND


class OCRUnavailable(RuntimeError):
    pass


# 2. 공통 함수: 추출 텍스트 길이 제한
def limit_text(text):
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS]
        text += "\n[텍스트 길이 제한으로 일부 생략]"
    return text

# 3. TXT, PY 등의 텍스트 파일 읽기
def extract_plain_text(file_path):
    return file_path.read_text(
        encoding="utf-8",
        errors="replace"
    )

# 4. Jupyter Notebook(.ipynb) 파일 읽기
def extract_ipynb(file_path):
    with file_path.open("r", encoding="utf-8") as f:
        notebook = json.load(f)

    texts = []

    for cell_num, cell in enumerate(notebook.get("cells", []), start=1):
        cell_type = cell.get("cell_type", "")
        source = cell.get("source", "")

        # source가 문자열 또는 문자열 리스트인 경우 모두 처리
        if isinstance(source, list):
            source = "".join(source)

        # 코드 셀과 마크다운 셀의 원본 텍스트 추출
        if source.strip():
            if cell_type == "code":
                texts.append(f"[코드 셀 {cell_num}]\n{source}")

            elif cell_type == "markdown":
                texts.append(f"[설명 셀 {cell_num}]\n{source}")

        # 코드 셀의 실행 결과(outputs) 추출
        if cell_type == "code":
            for output in cell.get("outputs", []):
                output_type = output.get("output_type", "")

                # 일반 텍스트 출력 및 print() 결과
                if output_type == "stream":
                    output_text = output.get("text", [])
                    if isinstance(output_text, list):
                        output_text = "".join(output_text)

                    if output_text.strip():
                        texts.append(
                            f"[실행 결과 셀 {cell_num}]\n{output_text}"
                        )

                # 실행 결과가 텍스트 또는 HTML 표로 저장된 경우
                elif output_type in ("execute_result", "display_data"):
                    data = output.get("data", {})

                    # 일반 텍스트 표현이 있으면 우선 추출
                    output_text = data.get("text/plain", [])
                    if isinstance(output_text, list):
                        output_text = "".join(output_text)

                    if output_text.strip():
                        texts.append(
                            f"[실행 결과 셀 {cell_num}]\n{output_text}"
                        )

                    # text/plain이 없으면 HTML 표현을 추출
                    elif "text/html" in data:
                        html_text = data["text/html"]
                        if isinstance(html_text, list):
                            html_text = "".join(html_text)

                        from bs4 import BeautifulSoup
                        plain_text = BeautifulSoup(
                            html_text, "html.parser"
                        ).get_text(" ", strip=True)

                        if plain_text:
                            texts.append(
                                f"[실행 결과 셀 {cell_num}]\n{plain_text}"
                            )

    return "\n\n".join(texts)


# 5. CSV, TSV 파일 읽기
def extract_csv(file_path):
    last_error = None

    # TSV는 탭, CSV는 쉼표로 구분
    sep = "\t" if file_path.suffix.lower() == ".tsv" else ","

    for encoding in ["utf-8-sig", "cp949"]:
        try:
            df = pd.read_csv(
                file_path,
                encoding=encoding,
                sep=sep,
                nrows=MAX_TABLE_ROWS
            )

            text = "컬럼: " + ", ".join(map(str, df.columns))
            text += "\n\n" + df.to_string(index=False)

            return text

        except (UnicodeDecodeError, pd.errors.ParserError) as e:
            last_error = e

    raise ValueError(f"CSV/TSV 읽기 실패: {last_error}")

# 6. XLSX 파일 읽기
def extract_xlsx(file_path):
    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True
    )

    texts = []

    try:
        for sheet in workbook.worksheets:
            texts.append(f"[시트명: {sheet.title}]")

            for row_num, row in enumerate(
                sheet.iter_rows(values_only=True), start=1
            ):
                if row_num > MAX_TABLE_ROWS:
                    texts.append("[행 제한으로 일부 생략]")
                    break

                values = [
                    str(value) if value is not None else ""
                    for value in row
                ]

                if any(values):
                    texts.append(" | ".join(values))

    finally:
        workbook.close()

    return "\n".join(texts)

# 7. PDF 파일 읽기 (OCR 지원)
def extract_pdf(file_path):
    texts = []

    with pymupdf.open(file_path) as pdf:
        for page_num, page in enumerate(pdf):
            if page_num >= MAX_PDF_PAGES:
                texts.append("[페이지 제한으로 일부 생략]")
                break

            # 일반 텍스트 추출
            text = page.get_text()

            # 텍스트가 없으면 OCR 실행
            if not text.strip():
                if pytesseract is None or not OCR_COMMAND:
                    raise OCRUnavailable("Tesseract와 kor/eng 언어 데이터가 필요합니다.")
                print(f"  OCR 처리 중: {page_num + 1}페이지")

                pix = page.get_pixmap(
                    matrix=pymupdf.Matrix(2, 2),
                    alpha=False
                )

                image = Image.frombytes(
                    "RGB",
                    (pix.width, pix.height),
                    pix.samples
                )

                text = pytesseract.image_to_string(
                    image,
                    lang="kor+eng"
                )

            if text.strip():
                texts.append(f"[페이지 {page_num + 1}]\n{text}")

    return "\n\n".join(texts)


# 8. DOCX 파일 읽기
def extract_docx(file_path):
    doc = Document(file_path)
    texts = []

    # 본문 문단
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            texts.append(paragraph.text)

    # 표 안의 내용
    for table in doc.tables:
        for row in table.rows:
            texts.append(
                " | ".join(cell.text for cell in row.cells)
            )

    return "\n".join(texts)

# 9. PPTX 파일 읽기
def extract_pptx(file_path):
    presentation = Presentation(file_path)
    texts = []

    for slide_num, slide in enumerate(presentation.slides, start=1):
        texts.append(f"[슬라이드 {slide_num}]")

        for shape in slide.shapes:

            # 일반 텍스트
            if shape.has_text_frame and shape.text.strip():
                texts.append(shape.text)

            # 표
            if shape.has_table:
                for row in shape.table.rows:
                    texts.append(
                        " | ".join(cell.text for cell in row.cells)
                    )

    return "\n".join(texts)

# 10. 파일 확장자에 따라 추출 함수 선택
def extract_content(file_path):
    if not file_path.exists():
        return "", "file_not_found"

    ext = file_path.suffix.lower()

    try:
        if ext in {
            ".py", ".txt", ".md", ".json",
            ".html", ".css", ".js", ".sql", ".r"
        }:
            text = extract_plain_text(file_path)

        elif ext == ".ipynb":
            text = extract_ipynb(file_path)

        elif ext in {".csv", ".tsv"}:
            text = extract_csv(file_path)

        elif ext == ".xlsx":
            text = extract_xlsx(file_path)

        elif ext == ".pdf":
            text = extract_pdf(file_path)

        elif ext == ".docx":
            text = extract_docx(file_path)

        elif ext == ".pptx":
            text = extract_pptx(file_path)

        else:
            return "", "unsupported_extension"

        text = limit_text(text)

        if not text.strip():
            return "", "empty_text"

        return text, "success"

    except Exception as e:
        return "", f"error: {type(e).__name__}: {e}"


# 11. 기존 CSV 읽기
metadata = pd.read_csv(
    METADATA_CSV,
    encoding="utf-8-sig"
)

# 필수 열이 존재하는지 확인
required_columns = {"file_id", "path"}

if not required_columns.issubset(metadata.columns):
    raise ValueError(
        f"CSV에 다음 열이 필요해: {required_columns}\n"
        f"현재 열: {list(metadata.columns)}"
    )

# Reuse previous successful OCR only when the original file hash is unchanged.
previous_texts, previous_audit = {}, {}
for name, target in ((".previous_extracted_text.csv", previous_texts),
                     (".previous_audit.csv", previous_audit)):
    candidate = BASE_DIR / name
    if candidate.is_file():
        frame = pd.read_csv(candidate, dtype=str, keep_default_na=False, encoding="utf-8-sig")
        target.update({row["file_id"]: row for row in frame.to_dict("records")})

# 12. 각 파일의 내용 추출
results = []

for _, row in metadata.iterrows():

    file_id = row["file_id"]
    file_path = resolve_path(row["path"])

    print(f"처리 중: {file_path.name}")

    text, status = extract_content(file_path)
    if status.startswith("error: OCRUnavailable"):
        old = previous_texts.get(str(file_id), {})
        audit = previous_audit.get(str(file_id), {})
        if (old.get("extraction_status") == "success"
                and old.get("filename") == file_path.name
                and audit.get("source_sha256") == hashlib.sha256(file_path.read_bytes()).hexdigest()):
            text, status = old["extracted_text"], "success"
            print("  원본 해시가 동일한 기존 OCR 결과 재사용")

    results.append({
        "file_id": file_id,
        "filename": file_path.name,
        "extension": file_path.suffix.lower(),
        "path": str(row["path"]).replace("\\", "/"),
        "extracted_text": text,
        "char_count": len(text),
        "extraction_status": status
    })

# 13. 결과 저장
result_df = pd.DataFrame(results, columns=["file_id", "filename", "extension", "path", "extracted_text", "char_count", "extraction_status"])

# metadata.csv에만 있는 추가 정보 가져오기
extra_columns = [
    col for col in metadata.columns
    if col not in result_df.columns and col != "file_id"
]

# file_id를 기준으로 두 데이터 합치기
if extra_columns:
    result_df = result_df.merge(
        metadata[["file_id"] + extra_columns],
        on="file_id",
        how="left"
    )

# 14. 결과 저장
result_df.to_csv(
    OUTPUT_CSV,
    index=False,
    encoding="utf-8-sig"
)

print("\n텍스트 추출 및 파일 정보 결합 완료!")
print(f"총 파일 수: {len(result_df)}")
print(
    f"추출 성공: "
    f"{(result_df['extraction_status'] == 'success').sum()}"
)
print(f"저장 위치: {OUTPUT_CSV.resolve()}")

# 14. 결과 미리 보기
result_df["preview"] = (
    result_df["extracted_text"]
    .str.replace("\n", " ", regex=False)
    .str.slice(0, 100)
)

print(
    result_df[
        [
            "file_id",
            "filename",
            "char_count",
            "extraction_status",
            "preview"
        ]
    ].head(10).to_string(index=False)
)